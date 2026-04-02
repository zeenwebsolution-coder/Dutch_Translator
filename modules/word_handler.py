from __future__ import annotations

import io
from typing import Dict, List, Set

from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.text.run import Run


# ── Helpers ────────────────────────────────────────────────────────────────

def _clean_text(text: str) -> str:
    if text is None:
        return ""
    val = str(text).strip()
    if not val or val.lower() in ("none", "nan"):
        return ""
    return val


def _iter_block_items(doc: Document):
    """
    Yield paragraphs and tables in document order.
    """
    for block in doc.element.body:
        if block.tag.endswith('p'):
            yield Paragraph(block, doc)
        elif block.tag.endswith('tbl'):
            yield Table(block, doc)


def _extract_runs_from_paragraph(paragraph: Paragraph) -> List[Run]:
    return list(paragraph.runs)


# ── Reading ────────────────────────────────────────────────────────────────

def read_word_entries(file_bytes: bytes) -> Dict[str, List[str]]:
    """
    Extract all unique textual entries from a Word document.
    Returns:
        {
            "doc": [list of unique strings]
        }
    """
    doc = Document(io.BytesIO(file_bytes))

    seen: Set[str] = set()
    results: List[str] = []

    for block in _iter_block_items(doc):
        if isinstance(block, Paragraph):
            text = _clean_text(block.text)
            if text and text not in seen:
                seen.add(text)
                results.append(text)

        elif isinstance(block, Table):
            for row in block.rows:
                for cell in row.cells:
                    cell_text = _clean_text(cell.text)
                    if cell_text and cell_text not in seen:
                        seen.add(cell_text)
                        results.append(cell_text)

    return {"doc": results}


# ── Writing Translations ───────────────────────────────────────────────────

def write_translations(
    original_bytes: bytes,
    translation_cache: Dict[str, str],
) -> bytes:
    """
    Replace text in the Word document using translation_cache.

    - Matches full paragraph text OR individual runs
    - Preserves formatting as much as possible
    """
    doc = Document(io.BytesIO(original_bytes))

    for block in _iter_block_items(doc):

        if isinstance(block, Paragraph):
            _replace_in_paragraph(block, translation_cache)

        elif isinstance(block, Table):
            for row in block.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        _replace_in_paragraph(paragraph, translation_cache)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def _replace_in_paragraph(paragraph: Paragraph, translation_cache: Dict[str, str]):
    # Try full paragraph match first
    full_text = _clean_text(paragraph.text)
    if full_text and full_text in translation_cache:
        translated = translation_cache[full_text]
        # Put translation in first run, clear the rest
        if paragraph.runs:
            paragraph.runs[0].text = translated
            for run in paragraph.runs[1:]:
                run.text = ""
        return

    # Fallback: try run-by-run match
    for run in paragraph.runs:
        cleaned = _clean_text(run.text)
        if cleaned and cleaned in translation_cache:
            run.text = translation_cache[cleaned]
# ── Utility ────────────────────────────────────────────────────────────────

def unique_words(word_entries: Dict[str, List[str]]) -> List[str]:
    """
    Deduplicate extracted words.
    """
    seen: Set[str] = set()
    result: List[str] = []

    for entries in word_entries.values():
        for word in entries:
            if word not in seen:
                seen.add(word)
                result.append(word)

    return result